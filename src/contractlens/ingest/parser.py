"""Extract plain text from PDF, DOCX, TXT/MD files (path or bytes)."""
from __future__ import annotations

import io
import re
import unicodedata
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class UnsupportedFormat(ValueError):
    pass


def _pdf(data: bytes) -> str:
    import pdfplumber
    pages = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            pages.append(page.extract_text(x_tolerance=1.5, y_tolerance=3) or "")
    text = "\n\n".join(pages)
    if len(text.strip()) < 50:
        raise ValueError("PDF has no extractable text layer (scanned image?). Run OCR first.")
    return text


def _docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("­", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)          # de-hyphenate line breaks
    text = re.sub(r"\n[ \t]*\d{1,3}[ \t]*\n", "\n", text)  # bare page numbers
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(source: str | Path | bytes, filename: str | None = None) -> str:
    if isinstance(source, (str, Path)) and not isinstance(source, bytes):
        path = Path(source)
        filename = filename or path.name
        data = path.read_bytes()
    else:
        data = source
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED:
        raise UnsupportedFormat(f"Unsupported file type '{suffix}'. Supported: {sorted(SUPPORTED)}")
    if suffix == ".pdf":
        raw = _pdf(data)
    elif suffix == ".docx":
        raw = _docx(data)
    else:
        raw = data.decode("utf-8", errors="replace")
    return normalize(raw)
