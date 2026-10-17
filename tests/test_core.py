import io

import numpy as np
import pytest

from contractlens.analysis.analyzer import ContractAnalyzer
from contractlens.analysis.entities import extract_entities, extract_governing_law
from contractlens.analysis.risk import assess
from contractlens.config import CATEGORY_NAMES
from contractlens.data.cuad import document_labels, label_segments, to_segment_dataset
from contractlens.ingest.parser import UnsupportedFormat, extract_text, normalize
from contractlens.ingest.segment import segment
from contractlens.models.classifier import ClauseClassifier, evaluate, precision_at_recall
from contractlens.search.retriever import BM25

# ------------------------------------------------------------------ data


def test_official_split_is_disjoint(cuad):
    train, test = cuad
    assert train and test
    assert not {c.title for c in train} & {c.title for c in test}
    assert set(train[0].spans) == set(CATEGORY_NAMES)


def test_segments_keep_offsets_and_cover_all_spans(cuad):
    for c in cuad[0][:10]:
        segs = segment(c.text)
        assert all(0 <= s.start < s.end <= len(c.text) for s in segs)
        assert all(a.end <= b.start for a, b in zip(segs, segs[1:]))  # ordered, non-overlapping
        for spans in c.spans.values():
            for a0, a1, _ in spans:
                assert any(s.start < a1 and s.end > a0 for s in segs)


def test_segmenter_does_not_split_on_abbreviations():
    text = ("This Agreement is made between Acme Inc. and Beta Corp. as of the date below and shall remain in "
            "force for the term described in Sec. 4 hereof unless terminated earlier.")
    assert len(segment(text, max_chars=60)) <= 2


def test_label_segments_marks_overlap(cuad):
    c = next(c for c in cuad[0] if c.spans["Governing Law"])
    segs = segment(c.text)
    y = label_segments(c, segs)
    k = CATEGORY_NAMES.index("Governing Law")
    assert y[:, k].sum() >= 1
    assert "law" in " ".join(segs[i].text.lower() for i in np.where(y[:, k])[0])


# ----------------------------------------------------------------- model


def test_classifier_learns_and_evaluates(cuad, trained):
    clf, _ = trained
    _, test = cuad
    Xt, Yt, Dt = to_segment_dataset(test)
    m = evaluate(clf, Xt, Yt, Dt, document_labels(test))
    assert set(m["per_category"]) == set(CATEGORY_NAMES)
    # even a tiny model must clearly beat chance on frequent, lexical clauses
    assert m["per_category"]["Governing Law"]["aupr"] > 0.5
    assert m["summary"]["macro_aupr"] > 0.1


def test_save_load_roundtrip(trained):
    clf, path = trained
    loaded = ClauseClassifier.load(path)
    texts = ["This Agreement shall be governed by the laws of the State of New York.",
             "Neither party may assign this Agreement without prior written consent."]
    np.testing.assert_allclose(clf.predict_proba(texts), loaded.predict_proba(texts))


def test_precision_at_recall():
    y = np.array([1, 0, 1, 0, 1])
    p = np.array([0.9, 0.8, 0.7, 0.2, 0.1])
    assert precision_at_recall(y, p, 0.66) == pytest.approx(2 / 3)
    assert precision_at_recall(y, p, 1.0) == pytest.approx(0.6)


def test_top_terms_are_readable(trained):
    clf, _ = trained
    terms = [t for t, _ in clf.top_terms("Governing Law", 20)]
    assert any("law" in t or "governed" in t for t in terms)


# -------------------------------------------------------------- analysis


def test_entities():
    text = ("This Supply Agreement is entered into on March 3, 2019 by and between Acme Widgets, Inc., a "
            "Delaware corporation (\"Supplier\"), and Beta Retail LLC (\"Buyer\"). Buyer shall pay "
            "$250,000 within thirty (30) days. This Agreement shall be governed by the laws of the State "
            "of New York. A late fee of 1.5% per month applies.")
    e = extract_entities(text)
    assert e["governing_law"] == "New York"
    assert "March 3, 2019" in e["dates"]
    assert any("250,000" in m for m in e["monetary_amounts"])
    assert any("thirty" in d.lower() for d in e["durations"])
    assert "1.5%" in e["percentages"]
    assert e["parties"][0].startswith("Acme Widgets")
    assert extract_governing_law("no jurisdiction mentioned here") is None


def test_risk_rules():
    ev = [{"text": "x", "confidence": 0.9}]
    findings, score = assess({"Uncapped Liability": ev, "Renewal Term": ev, "Governing Law": ev,
                              "Expiration Date": ev, "Anti-Assignment": ev, "Insurance": ev})
    titles = {f.title for f in findings}
    assert "Uncapped liability" in titles
    assert "No limitation-of-liability clause found" in titles
    assert "Renewal without a clear opt-out notice period" in titles
    assert findings[0].severity == "HIGH"
    assert 0 < score <= 100
    _, low = assess({c: ev for c in ("Cap On Liability", "Governing Law", "Expiration Date",
                                     "Anti-Assignment", "Insurance")})
    assert low < score


def test_bm25_ranks_relevant_segment_first():
    docs = ["Payment is due within 30 days of invoice.",
            "Either party may terminate this Agreement upon ninety days written notice.",
            "This Agreement is governed by the laws of Texas."]
    assert BM25(docs).top_k("how can we terminate the contract?")[0][0] == 1


def test_analyzer_end_to_end(trained, sample_text):
    clf, _ = trained
    rep = ContractAnalyzer(clf).analyze(sample_text, "msa").to_dict()
    assert rep["n_segments"] > 20
    assert 0 <= rep["risk_score"] <= 100
    assert rep["risk_level"] in {"LOW", "MEDIUM", "HIGH"}
    assert rep["summary"]
    for hits in rep["clauses"].values():
        for h in hits:
            assert sample_text[h["start"]:h["end"]].split()[0] in h["text"]
    res = ContractAnalyzer(clf).search(sample_text, "termination notice", k=3)
    assert res and "terminat" in res[0]["text"].lower()


# ---------------------------------------------------------------- parser


def test_parser_txt_and_docx():
    assert "Agreement" in extract_text(b"This Agreement   is\r\nbinding.", filename="a.txt")
    import docx
    d = docx.Document()
    d.add_paragraph("MASTER SERVICES AGREEMENT")
    d.add_paragraph("The parties agree as follows.")
    buf = io.BytesIO()
    d.save(buf)
    assert "MASTER SERVICES AGREEMENT" in extract_text(buf.getvalue(), filename="c.docx")
    with pytest.raises(UnsupportedFormat):
        extract_text(b"x", filename="c.xls")


def test_normalize_dehyphenates():
    assert normalize("termi-\nnation  of the") == "termination of the"


def test_parties_from_real_preambles():
    from contractlens.analysis.entities import extract_parties
    assert extract_parties('BETWEEN: ORGANIC PREPARATIONS INC. 2nd Floor, Port Vila "the Manufacturer" -- AND -- '
                           'AGAPE ATP INTERNATIONAL HOLDING LIMITED Unit 05, Hong Kong "the Customer"\n') == \
        ["ORGANIC PREPARATIONS INC.", "AGAPE ATP INTERNATIONAL HOLDING LIMITED"]
    assert extract_parties('is between Clear Capital (the "Company") and RadialSpark, LLC (the "Contractor"). ') == \
        ["Clear Capital", "RadialSpark, LLC"]
